"""Excel report writer (openpyxl).

Produces a three-sheet workbook:

* **Jobs** - one formatted table row per vacancy, with clickable apply links
* **Summary** - run metadata plus company / location / work-mode / skill counts
* **Run Log** - UiPath-style transaction log with per-item status and timings
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .extractors import clean_text, parse_date
from .models import JobPosting

JOBS_COLUMNS: list[tuple[str, str, int]] = [
    ("Job ID", "job_id", 12),
    ("Job Title", "title", 42),
    ("Company", "company", 30),
    ("Location", "location", 24),
    ("Work Mode", "work_mode", 12),
    ("Employment Type", "employment_type", 16),
    ("Salary", "salary", 22),
    ("Posted Date", "posted_date", 12),
    ("Category", "category", 22),
    ("Skills", "skills", 38),
    ("Description", "description", 62),
    ("Apply URL", "apply_url", 42),
    ("Source URL", "source_url", 42),
    ("Portal", "portal", 14),
    ("Scraped At", "scraped_at", 22),
]

LOG_COLUMNS: list[tuple[str, str, int]] = [
    ("Timestamp", "timestamp", 22),
    ("Level", "level", 8),
    ("Status", "status", 18),
    ("Step", "step", 16),
    ("Job ID", "job_id", 12),
    ("Job Title", "title", 36),
    ("Attempt", "attempt", 8),
    ("Duration (ms)", "duration_ms", 13),
    ("Message", "message", 90),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1F3864", bold=True, size=16)
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
SECTION_FONT = Font(color="1F3864", bold=True, size=11)
MUTED_FONT = Font(color="7F7F7F", size=10)
LINK_FONT = Font(color="0563C1", underline="single")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))

STATUS_FILLS = {
    "Successful": PatternFill("solid", fgColor="E2EFDA"),
    "Skipped": PatternFill("solid", fgColor="F2F2F2"),
    "BusinessException": PatternFill("solid", fgColor="FFF2CC"),
    "SystemException": PatternFill("solid", fgColor="FBE2E2"),
}


# --------------------------------------------------------------------------- public
def write_report(result, config: dict) -> Path:
    """Write ``result.jobs`` and ``result.records`` to an .xlsx file."""
    destination = Path(result.output_path) if result.output_path else _default_destination(config)
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.properties.title = "Job Vacancy Report"
    workbook.properties.creator = "RPA Job Bot"
    workbook.properties.subject = f"Vacancies collected from {result.portal_label}"

    jobs_ws = workbook.active
    jobs_ws.title = "Jobs"
    _write_jobs_sheet(jobs_ws, result.jobs)

    summary_ws = workbook.create_sheet("Summary")
    _write_summary_sheet(summary_ws, result)

    log_ws = workbook.create_sheet("Run Log")
    _write_log_sheet(log_ws, result.records)

    workbook.save(destination)
    return destination


def read_existing_jobs(path: str | Path) -> list[JobPosting]:
    """Load the Jobs sheet of an earlier report (used by ``--append``)."""
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet = workbook["Jobs"] if "Jobs" in workbook.sheetnames else workbook.active

    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []

    index = {str(name).strip(): position for position, name in enumerate(header) if name}
    jobs: list[JobPosting] = []
    for row in rows:
        def value(key: str) -> str:
            position = index.get(key)
            if position is None or position >= len(row) or row[position] is None:
                return ""
            return clean_text(row[position])

        title = value("Job Title")
        if not title:
            continue

        posted = row[index["Posted Date"]] if "Posted Date" in index else None
        if isinstance(posted, (datetime, date)):
            posted_text = posted.strftime("%Y-%m-%d")
        else:
            posted_text = parse_date(posted)

        skills_raw = value("Skills")
        jobs.append(
            JobPosting(
                portal=value("Portal") or "unknown",
                title=title,
                company=value("Company"),
                location=value("Location"),
                work_mode=value("Work Mode") or "Unspecified",
                employment_type=value("Employment Type"),
                salary=value("Salary"),
                posted_date=posted_text,
                category=value("Category"),
                skills=[part.strip() for part in skills_raw.split(",") if part.strip()],
                description=value("Description"),
                apply_url=value("Apply URL"),
                source_url=value("Source URL"),
                scraped_at=value("Scraped At"),
            )
        )
    return jobs


# --------------------------------------------------------------------------- sheets
def _write_jobs_sheet(sheet, jobs: list[JobPosting]) -> None:
    headers = [header for header, _attr, _width in JOBS_COLUMNS]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for row_number, job in enumerate(jobs, start=2):
        data = job.as_dict()
        values: list[Any] = []
        for header, attr, _width in JOBS_COLUMNS:
            if attr == "posted_date":
                values.append(_as_excel_date(job.posted_date))
            elif attr == "skills":
                values.append(", ".join(job.skills))
            else:
                values.append(_safe_cell(data.get(attr, "")))

        sheet.append(values)

        for column, (_header, attr, _width) in enumerate(JOBS_COLUMNS, start=1):
            cell = sheet.cell(row=row_number, column=column)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=attr in {"description", "skills", "title"},
            )
            cell.border = THIN_BORDER
            if attr in {"apply_url", "source_url"} and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT
            if attr == "posted_date":
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center", vertical="top")

        sheet.row_dimensions[row_number].height = 46 if job.description else 20

    for position, (_header, _attr, width) in enumerate(JOBS_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width

    sheet.freeze_panes = "A2"
    last_row = max(len(jobs) + 1, 2)
    table = Table(
        displayName="JobsTable",
        ref=f"A1:{get_column_letter(len(JOBS_COLUMNS))}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"


def _write_summary_sheet(sheet, result) -> None:
    sheet.sheet_view.showGridLines = False
    widths = {"A": 24, "B": 38, "C": 2, "D": 30, "E": 10, "F": 2, "G": 30, "H": 10}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:H1")
    title = sheet["A1"]
    title.value = "Job Vacancy Collection — Run Summary"
    title.font = TITLE_FONT
    sheet.row_dimensions[1].height = 26

    sheet["A2"] = f"Generated by RPA Job Bot · {result.portal_label}"
    sheet["A2"].font = MUTED_FONT
    sheet.merge_cells("A2:H2")

    run_info = [
        ("Run started", _friendly_timestamp(result.started_at)),
        ("Run finished", _friendly_timestamp(result.finished_at)),
        ("Duration", f"{result.duration_seconds:.1f} s"),
        ("Portal", result.portal),
        ("Target board", result.portal_label),
        ("Report file", Path(result.output_path).name if result.output_path else ""),
    ]
    if getattr(result, "keywords", None):
        run_info.append(("Keyword filter", ", ".join(result.keywords)))
    if getattr(result, "countries", None):
        run_info.append(("Country filter", ", ".join(result.countries)))
    cursor = 4
    cursor = _write_block(sheet, 1, cursor, "Run information", run_info, "Item", "Value")

    results = [
        ("Vacancies discovered", result.discovered),
        ("New vacancies this run", result.successful),
        ("Duplicates skipped", result.skipped),
        ("Outside location filter", getattr(result, "filtered_location", 0)),
        ("Filtered out by keyword", getattr(result, "filtered", 0)),
        ("Failed transactions", result.failed),
        ("Existing rows kept", result.carried_over),
        ("Total rows in report", len(result.jobs)),
    ]
    if not getattr(result, "filtered_location", 0):
        results = [row for row in results if row[0] != "Outside location filter"]
    if not getattr(result, "filtered", 0):
        results = [row for row in results if row[0] != "Filtered out by keyword"]
    cursor = _write_block(sheet, 1, cursor + 1, "Results", results, "Metric", "Count")

    sources = _counter_rows((job.portal for job in result.jobs), 10)
    source_cursor = _write_block(
        sheet, 1, cursor + 1, "Jobs by source", sources, "Source", "Jobs"
    )

    companies = _counter_rows((job.company for job in result.jobs), 10)
    locations = _counter_rows((job.location for job in result.jobs), 10)
    work_modes = _counter_rows((job.work_mode for job in result.jobs), 6)
    skills = _counter_rows(
        (skill for job in result.jobs for skill in job.skills), 12, preserve_case=True
    )

    right_column = 4
    _write_block(sheet, right_column, 4, "Top companies", companies, "Company", "Jobs")
    _write_block(sheet, right_column + 3, 4, "Top locations", locations, "Location", "Jobs")
    _write_block(sheet, right_column, 16, "Work mode", work_modes, "Mode", "Jobs")
    _write_block(sheet, right_column + 3, 16, "Top skills", skills, "Skill", "Mentions")

    note_row = max(source_cursor + 1, 34)
    sheet.cell(row=note_row, column=1, value="Data sources & compliance").font = SECTION_FONT
    sheet.cell(row=note_row + 1, column=1, value=_source_note(result)).font = MUTED_FONT
    sheet.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=8)
    sheet.cell(row=note_row + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[note_row + 1].height = 48

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _source_note(result) -> str:
    from .portals import PORTALS

    labels = []
    for name in sorted({job.portal for job in result.jobs}):
        portal_cls = PORTALS.get(name)
        labels.append(portal_cls.label if portal_cls else name)
    credit = ", ".join(labels) if labels else result.portal_label
    return (
        f"Data sources: {credit}. Source credit is preserved in the Portal, Apply URL and "
        "Source URL columns - keep those links when sharing the workbook. The bot respects "
        "robots.txt, throttles requests per host, retries transient failures and stores only "
        "publicly listed vacancy data."
    )


def _write_log_sheet(sheet, records) -> None:
    headers = [header for header, _attr, _width in LOG_COLUMNS]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 22

    for row_number, record in enumerate(records, start=2):
        data = record.as_dict()
        sheet.append([_safe_cell(data[attr]) for _header, attr, _width in LOG_COLUMNS])

        fill = STATUS_FILLS.get(record.status)
        for column in range(1, len(LOG_COLUMNS) + 1):
            cell = sheet.cell(row=row_number, column=column)
            cell.alignment = Alignment(vertical="top", wrap_text=column == len(LOG_COLUMNS))
            if fill is not None:
                cell.fill = fill

    for position, (_header, _attr, width) in enumerate(LOG_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LOG_COLUMNS))}{max(len(records) + 1, 2)}"


# --------------------------------------------------------------------------- helpers
def _write_block(sheet, start_col: int, start_row: int, title: str, pairs, key_label: str, value_label: str) -> int:
    """Render a small two-column block and return the next free row."""
    key_letter = get_column_letter(start_col)
    value_letter = get_column_letter(start_col + 1)

    header = sheet[f"{key_letter}{start_row}"]
    header.value = title
    header.font = SECTION_FONT
    header.fill = SECTION_FILL
    sheet.merge_cells(f"{key_letter}{start_row}:{value_letter}{start_row}")

    sheet[f"{key_letter}{start_row + 1}"] = key_label
    sheet[f"{value_letter}{start_row + 1}"] = value_label
    for cell in (sheet[f"{key_letter}{start_row + 1}"], sheet[f"{value_letter}{start_row + 1}"]):
        cell.font = Font(bold=True, color="595959", size=10)
        cell.border = Border(bottom=Side(style="thin", color="BFBFBF"))

    rows = list(pairs)
    if not rows:
        rows = [("No data", "—")]

    for offset, (key, value) in enumerate(rows, start=2):
        sheet[f"{key_letter}{start_row + offset}"] = _safe_cell(key)
        sheet[f"{value_letter}{start_row + offset}"] = value

    return start_row + len(rows) + 3


def _counter_rows(values: Iterable[str], limit: int, preserve_case: bool = False) -> list[tuple[str, int]]:
    counter: Counter[str] = Counter()
    for value in values:
        cleaned = clean_text(value)
        if cleaned and cleaned.casefold() not in {"unspecified", "unknown", "n/a", "—"}:
            counter[cleaned] += 1
    return counter.most_common(limit)


def _as_excel_date(value: str):
    if not value:
        return ""
    try:
        return date.fromisoformat(value)
    except ValueError:
        return value


def _friendly_timestamp(value: str) -> str:
    if not value:
        return ""
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M:%S UTC")
    except ValueError:
        return value


def _safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection from scraped text."""
    if isinstance(value, str) and value[:1] in {"=", "+", "@"}:
        return "'" + value
    return value


def _default_destination(config: dict) -> Path:
    from .config import output_path

    return output_path(config)
