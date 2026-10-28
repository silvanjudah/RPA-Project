"""End-to-end bot tests with a stubbed HTTP layer (no network required)."""

from __future__ import annotations

import copy
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rpa_jobs.config import load_config
from rpa_jobs.pipeline import JobBot, matches_keywords
from rpa_jobs.models import JobPosting
from rpa_jobs.report import _safe_cell, read_existing_jobs

from .test_portals import StubHttp

BASE = "https://www.python.org/jobs/"


def card_html(title: str, company: str, location: str, href: str, posted: str = "2024-05-01") -> str:
    """One python.org job-board card (the markup the live adapter parses)."""
    return f"""
    <li>
      <h2 class="listing-company">
        <span class="listing-company-name"><a href="{href}">{title}</a><br>{company}</span>
        <span class="listing-location"><a href="#">{location}</a></span>
      </h2>
      <span class="listing-job-type"><a href="#">Full time</a></span>
      <span class="listing-posted">Posted: <time datetime="{posted}">{posted}</time></span>
      <span class="listing-company-category"><a href="#">Developer / Engineer</a></span>
    </li>
    """


def listing_html(*cards: str) -> str:
    return (
        '<html><body><ol class="list-recent-jobs">'
        + "".join(cards)
        + "</ol></body></html>"
    )


def detail_html(title: str, company: str, location: str, description: str) -> str:
    return f"""
    <html><body><div class="job-description">
      <p>{description}</p>
      <p>Location: {location}</p>
      <p><a href="https://jobs.acme.com/apply">Apply for this job</a></p>
    </div></body></html>
    """


def make_config(tmp_path: Path, *, fetch_details: bool = False, retries: int = 1) -> dict:
    config = copy.deepcopy(load_config())
    config["portal"] = "python-jobs"
    config["portals"]["python-jobs"]["base_url"] = BASE
    config["portals"]["python-jobs"]["max_pages"] = 1
    config["crawl"].update(
        {
            "max_items": 10,
            "fetch_details": fetch_details,
            "retries": retries,
            "backoff_seconds": 0.01,
            "delay_seconds": 0,
        }
    )
    config["output"]["directory"] = str(tmp_path)
    config["filters"] = {"keywords": [], "countries": [], "match": "any"}
    return config


def three_cards() -> str:
    return listing_html(
        card_html("Senior Python Developer", "Acme", "Remote (UK)", f"{BASE}1/"),
        card_html("Data Engineer", "Globex", "London, UK", f"{BASE}2/"),
        card_html("Backend Developer", "Initech", "Berlin, Germany", f"{BASE}3/"),
    )


def test_full_run_writes_formatted_workbook(tmp_path):
    stub = StubHttp({BASE: three_cards()})
    config = make_config(tmp_path)
    report_path = tmp_path / "report.xlsx"

    result = JobBot(config, http=stub, output=report_path).run()

    assert result.discovered == 3
    assert result.successful == 3
    assert result.failed == 0
    assert result.output_path == report_path
    assert report_path.exists()

    workbook = load_workbook(report_path)
    assert workbook.sheetnames == ["Jobs", "Summary", "Run Log"]

    jobs = workbook["Jobs"]
    assert jobs["A1"].value == "Job ID"
    assert jobs.freeze_panes == "A2"
    assert jobs.max_row == 4
    assert "JobsTable" in jobs.tables

    headers = {cell.value: cell.column for cell in jobs[1]}
    titles = {jobs.cell(row=row, column=headers["Job Title"]).value for row in range(2, 5)}
    assert titles == {"Senior Python Developer", "Data Engineer", "Backend Developer"}

    first_data_row = 2
    location = jobs.cell(row=first_data_row, column=headers["Location"]).value
    assert location
    assert jobs.cell(row=first_data_row, column=headers["Work Mode"]).value in {
        "Remote",
        "On-site",
        "Unspecified",
    }
    apply_cell = jobs.cell(row=first_data_row, column=headers["Apply URL"])
    assert apply_cell.hyperlink is not None
    posted_cell = jobs.cell(row=first_data_row, column=headers["Posted Date"])
    assert isinstance(posted_cell.value, date)
    assert posted_cell.number_format == "yyyy-mm-dd"

    summary = workbook["Summary"]
    assert "Run Summary" in summary["A1"].value
    assert summary["B6"].value  # run start timestamp
    assert "Total rows in report" in [summary.cell(row=row, column=1).value for row in range(1, 30)]

    run_log = workbook["Run Log"]
    assert run_log["A1"].value == "Timestamp"
    assert run_log.max_row > 5


def test_duplicate_vacancies_are_skipped(tmp_path):
    duplicate_listing = listing_html(
        card_html("Python Developer", "Acme", "London", "https://www.python.org/jobs/dup-a/"),
        card_html("Python Developer", "Acme", "London", "https://www.python.org/jobs/dup-b/"),
    )
    result = JobBot(
        make_config(tmp_path),
        http=StubHttp({BASE: duplicate_listing}),
        output=tmp_path / "dup.xlsx",
    ).run()

    assert result.discovered == 2
    assert result.successful == 1
    assert result.skipped == 1
    assert len(result.jobs) == 1
    assert any(record.status == "Skipped" for record in result.records)


def test_transient_detail_failure_is_retried(tmp_path):
    href = "https://www.python.org/jobs/retry/"
    detail = detail_html(
        "Retry Engineer",
        "Acme",
        "Remote",
        "Exciting Python and Django work. Full-time.",
    )
    stub = StubHttp(
        {BASE: listing_html(card_html("Retry Engineer", "Acme", "Remote", href)), href: detail},
        fail_times={href: 2},
    )
    config = make_config(tmp_path, fetch_details=True, retries=3)
    result = JobBot(config, http=stub, output=tmp_path / "retry.xlsx").run()

    assert result.successful == 1
    assert result.failed == 0
    assert result.jobs[0].description.startswith("Exciting Python")
    assert stub.calls.count(href) == 3
    statuses = [record.status for record in result.records]
    assert statuses.count("SystemException") >= 2
    assert "Successful" in statuses


def test_permanent_detail_failure_is_marked_as_failed(tmp_path):
    href = "https://www.python.org/jobs/broken/"
    stub = StubHttp(
        {BASE: listing_html(card_html("Broken Job", "Acme", "Remote", href)), href: detail_html("x", "y", "z", "d")},
        always_fail={href},
    )
    config = make_config(tmp_path, fetch_details=True, retries=2)
    result = JobBot(config, http=stub, output=tmp_path / "failed.xlsx").run()

    assert result.discovered == 1
    assert result.successful == 0
    assert result.failed == 1
    assert result.jobs == []
    assert (tmp_path / "failed.xlsx").exists()
    assert any(record.status == "SystemException" and record.level == "ERROR" for record in result.records)


def test_append_mode_keeps_existing_rows_and_skips_duplicates(tmp_path):
    report_path = tmp_path / "weekly.xlsx"
    stub = StubHttp({BASE: three_cards()})

    first = JobBot(make_config(tmp_path), http=stub, output=report_path).run()
    assert first.successful == 3

    second = JobBot(make_config(tmp_path), http=stub, output=report_path, append=True).run()
    assert second.successful == 0
    assert second.carried_over == 3
    assert len(second.jobs) == 3

    jobs = load_workbook(report_path)["Jobs"]
    assert jobs.max_row == 4  # header + 3, no duplicates
    assert len(read_existing_jobs(report_path)) == 3


def test_append_mode_survives_corrupt_existing_file(tmp_path):
    report_path = tmp_path / "corrupt.xlsx"
    report_path.write_text("this is not a workbook", encoding="utf-8")

    result = JobBot(
        make_config(tmp_path),
        http=StubHttp({BASE: three_cards()}),
        output=report_path,
        append=True,
    ).run()

    assert result.successful == 3
    assert result.carried_over == 0
    assert load_workbook(report_path)["Jobs"].max_row == 4


def test_unknown_portal_configuration_is_rejected(tmp_path):
    config = make_config(tmp_path)
    config["portal"] = "does-not-exist"
    from rpa_jobs.portals import BusinessError

    with pytest.raises(BusinessError):
        JobBot(config, http=StubHttp({})).run()


def test_missing_location_is_marked_not_specified(tmp_path):
    listing = listing_html(
        card_html("Remote Role", "Acme", "", "https://www.python.org/jobs/remote/")
    )
    result = JobBot(
        make_config(tmp_path),
        http=StubHttp({BASE: listing}),
        output=tmp_path / "remote.xlsx",
    ).run()

    assert result.jobs[0].location == "Not specified"


def test_country_filter_keeps_only_india_before_enrichment(tmp_path, monkeypatch):
    # The live python.org adapter also filters its own listings by country.
    # Disable that layer here so the pipeline's pre-enrichment filter is what
    # this test exercises (both layers share rpa_jobs.filters).
    from rpa_jobs.portals.python_jobs import PythonJobsPortal

    monkeypatch.setattr(PythonJobsPortal, "country_filter", lambda self: [])
    listing = listing_html(
        card_html("Backend Engineer", "Acme", "Bangalore, India", "https://www.python.org/jobs/india/"),
        card_html("Data Analyst", "Globex", "Indiana, USA", "https://www.python.org/jobs/usa/"),
        card_html("Remote Engineer", "Initech", "Remote", "https://www.python.org/jobs/worldwide/"),
    )
    config = make_config(tmp_path)
    config["filters"]["countries"] = ["India"]
    result = JobBot(
        config,
        http=StubHttp({BASE: listing}),
        output=tmp_path / "india.xlsx",
    ).run()

    assert result.discovered == 3
    assert result.filtered_location == 2
    assert result.successful == 1
    assert result.jobs[0].title == "Backend Engineer"
    assert result.countries == ["India"]
    assert any(record.step == "FILTER" for record in result.records)


def test_keyword_filter_keeps_only_matching_rows(tmp_path):
    config = make_config(tmp_path)
    config["filters"] = {"keywords": ["python"], "match": "any"}
    result = JobBot(
        config,
        http=StubHttp({BASE: three_cards()}),
        output=tmp_path / "keyword.xlsx",
    ).run()

    assert result.discovered == 3
    assert result.successful == 1
    assert result.filtered == 2
    assert [job.title for job in result.jobs] == ["Senior Python Developer"]
    assert any(record.step == "FILTER" and record.status == "Skipped" for record in result.records)


def test_keyword_filter_all_mode_requires_every_word(tmp_path):
    config = make_config(tmp_path)
    config["filters"] = {"keywords": ["python", "berlin"], "match": "all"}
    result = JobBot(
        config,
        http=StubHttp({BASE: three_cards()}),
        output=tmp_path / "keyword-all.xlsx",
    ).run()

    assert result.successful == 0
    assert result.filtered == 3
    assert result.jobs == []
    assert (tmp_path / "keyword-all.xlsx").exists()


def test_matches_keywords_covers_all_searchable_fields():
    job = JobPosting(
        portal="test",
        title="Backend Engineer",
        company="Acme",
        location="Berlin",
        description="We use Django and PostgreSQL.",
        skills=["Python", "Django"],
    )
    assert matches_keywords(job, ["python"])
    assert matches_keywords(job, ["acme", "django"])
    assert not matches_keywords(job, ["acme", "rust"], match_all=True)
    assert matches_keywords(job, ["acme", "django", "python"], match_all=True)


def test_formula_injection_is_neutralised():
    assert _safe_cell("=HYPERLINK('https://evil.invalid')") == "'=HYPERLINK('https://evil.invalid')"
    assert _safe_cell("@SUM(A1)").startswith("'")
    assert _safe_cell("Senior Python Developer") == "Senior Python Developer"
    assert _safe_cell(42) == 42
